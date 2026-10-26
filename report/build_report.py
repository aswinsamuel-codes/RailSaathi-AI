"""
Builds the project report (DOCX) from the results/ and figures/ folders.

Formatting: A4, 1-inch margins, Times New Roman 12 pt body, 1.5 line
spacing, justified text, bold 16/14 pt headings, page numbers bottom
centre, numbered figure/table captions and IEEE-style references.
"""
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG, RES = ROOT / "figures", ROOT / "results"
OUT = ROOT / "report" / "RailSaathi_AI_Report.docx"

STUDENT = "Aswin Samuel.A"
REGNO = "URK24CS9014"
GITHUB = "https://github.com/aswinsamuel-codes/RailSaathi-AI"

R = lambda name: pd.read_csv(RES / name, index_col=0)  # noqa: E731
J = lambda name: json.loads((RES / f"{name}.json").read_text())  # noqa: E731
pct = lambda v, d=1: f"{100 * float(v):.{d}f}%"  # noqa: E731


# ---------------------------------------------------------------------------
# low-level helpers
# ---------------------------------------------------------------------------
def set_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, ""), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            if text == "":
                el = OxmlElement("w:t")
                el.text = "Right-click and choose Update Field to build the table of contents."
        run._r.append(el)
    set_font(run, 12)
    return run


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        self._styles()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))

    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = "Times New Roman"
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        normal.font.size = Pt(12)
        pf = normal.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
            h = st[name]
            h.font.name = "Times New Roman"
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            h.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
            h.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.line_spacing = 1.5
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            h.paragraph_format.keep_with_next = True

    # -- content -----------------------------------------------------------
    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def para(self, text, bold_lead=None, align=None):
        p = self.doc.add_paragraph()
        if bold_lead:
            set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text))
        if align:
            p.alignment = align
        return p

    def bullets(self, items):
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if isinstance(it, tuple):
                set_font(p.add_run(it[0]), bold=True)
                set_font(p.add_run(it[1]))
            else:
                set_font(p.add_run(it))

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(2)
            set_font(p.add_run(f"{i}.  "), bold=True)
            set_font(p.add_run(it))

    def figure(self, name, caption, width=6.1):
        self.fig_no += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(FIG / f"{name}.png"), width=Inches(width))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.line_spacing = 1.15
        c.paragraph_format.space_after = Pt(10)
        set_font(c.add_run(f"Fig. {self.fig_no}. "), size=11, bold=True)
        set_font(c.add_run(caption), size=11)
        return self.fig_no

    def table(self, df, caption, col_widths=None, font=10, index=True):
        self.tab_no += 1
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.line_spacing = 1.15
        c.paragraph_format.space_before = Pt(6)
        set_font(c.add_run(f"Table {self.tab_no}. "), size=11, bold=True)
        set_font(c.add_run(caption), size=11)
        data = df.reset_index() if index else df
        t = self.doc.add_table(rows=1, cols=len(data.columns))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, col in enumerate(data.columns):
            cell = t.rows[0].cells[j]
            cell.text = ""
            set_font(cell.paragraphs[0].add_run(str(col)), size=font, bold=True)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "D9E2F3")
            cell._tc.get_or_add_tcPr().append(shd)
        for _, row in data.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = ""
                bold = isinstance(v, str) and v.endswith("*")
                set_font(cells[j].paragraphs[0].add_run(str(v)), size=font, bold=bold)
                cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        for row in t.rows:
            for j, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(1)
                if col_widths:
                    cell.width = Inches(col_widths[j])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return self.tab_no

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
def build():
    rp = Report()
    doc = rp.doc

    # ---------------- cover page --------------------------------------------
    for text, size, bold, space in [
        ("KARUNYA INSTITUTE OF TECHNOLOGY AND SCIENCES", 16, True, 4),
        ("Department of Computer Science and Engineering", 13, False, 30),
        ("23CS2024 – Data Science Ecosystem", 13, True, 4),
        ("Project Based Learning – 2026-2027 Odd Semester", 12, False, 40),
        ("RailSaathi AI", 22, True, 6),
        ("AI-Powered Complaint Management for Rail Madad:", 16, True, 2),
        ("Multimodal Categorisation, Cost-Sensitive Prioritisation", 15, True, 2),
        ("and Predictive Maintenance Analytics", 15, True, 24),
        ("Smart India Hackathon 2024 – Problem Statement SIH1711", 12, False, 2),
        ("Ministry of Railways", 12, False, 50),
        ("Submitted by", 12, False, 2),
        (STUDENT, 14, True, 2),
        (f"Register Number: {REGNO}", 12, False, 2),
        ("III B.Tech, Semester V", 12, False, 40),
        ("Course Coordinator: Dr. Naveen Sundar", 12, False, 2),
        ("October 2026", 12, False, 2),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space)
        p.paragraph_format.line_spacing = 1.15
        set_font(p.add_run(text), size=size, bold=bold)
    rp.page_break()

    # ---------------- table of contents ---------------------------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("TABLE OF CONTENTS"), size=16, bold=True)
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')

    # ---------------- body (new section with page numbers) -------------------
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.footer.is_linked_to_previous = False
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    body._sectPr.append(pg)
    fp = body.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE")
    for r in fp.runs:
        set_font(r, 11)
    # the first section (cover + contents) carries no page number
    doc.sections[0].footer.paragraphs[0].text = ""

    ov = R("dataset_overview.csv")["value"]
    comp = R("categorisation_model_comparison.csv")
    abl = R("categorisation_ablation.csv")
    rob = R("categorisation_robustness.csv")
    urg = R("urgency_results.csv")
    queue = pd.read_csv(RES / "queue_simulation.csv")
    mf = R("maintenance_forecast.csv")
    sla = R("sla_breach_model.csv")
    neg = R("negative_sentiment_by_department.csv")
    ext = pd.read_csv(RES / "entity_extraction.csv")
    dup = pd.read_csv(RES / "duplicate_grouping.csv")
    routing, hot, sent, gps, qs, mon = (J("routing_summary"), J("hotspot_detection"), J("sentiment_results"),
                                        J("gps_metadata"), J("queue_summary"), J("monitoring_summary"))
    best = comp.loc["Linear SVM (calibrated)"]
    prop = urg.loc["LR + minimum expected cost (proposed)"]
    lr_u = urg.loc["Logistic Regression (arg-max)"]
    qp = queue.pivot(index="Severity", columns="Policy", values="Mean wait (min)")
    dup_best = dup[(dup["method"] == "text + coach")].sort_values("pair_F1").iloc[-1]
    dup_txt = dup[(dup["method"] == "text only")].sort_values("pair_F1").iloc[-1]

    # 1 ---------------------------------------------------------------------
    rp.h1("1. Problem Statement")
    rp.para("Rail Madad is the single-window grievance platform of Indian Railways. Passengers lodge complaints "
            "through a mobile app, a web portal, the 139 helpline and social media, often attaching photographs, "
            "videos or voice notes. Each complaint receives a reference number, is assigned a category and is "
            "forwarded to the department that must act on it [1]. Problem statement SIH1711 of Smart India "
            "Hackathon 2024 [2] observes that this registration and categorisation is largely manual, which delays "
            "resolution, and that it is especially difficult when the passenger shares only an image, a video or "
            "an audio clip. The statement asks for AI that categorises and prioritises complaints, extracts extra "
            "information (OCR, image metadata), routes complaints to the right official, predicts recurring "
            "problems for proactive maintenance and monitors the process using feedback sentiment.")
    rp.para("In data-science terms this is a group of linked problems: multi-class text classification on noisy, "
            "code-mixed and multimodal input; ordinal severity prediction in which errors have unequal cost; "
            "information extraction; near-duplicate detection; count time-series forecasting; and sentiment "
            "and service-level analytics. As required by the course, the project focuses on these analytical "
            "methods rather than on a user interface.")

    # 2 ---------------------------------------------------------------------
    rp.h1("2. Objectives")
    rp.numbered([
        "Build a realistic, multimodal complaint dataset that reflects Rail Madad categories, departments, "
        "severity levels, language noise and operational seasonality.",
        "Categorise complaints into 13 classes and 7 departments by fusing text, OCR text and image tags, and "
        "measure what each input stream contributes.",
        "Route complaints automatically only when the model is confident, and defer the rest to a human desk.",
        "Predict urgency with a cost-sensitive decision rule that minimises dangerous under-prioritisation, and "
        "quantify the effect on waiting times through queue simulation.",
        "Extract train number, coach, PNR and location from text, OCR and EXIF/GPS metadata, and merge duplicate "
        "reports of the same incident into one ticket.",
        "Forecast weekly maintenance-related complaints per train and flag deteriorating trains early.",
        "Analyse feedback sentiment, predict service-level (SLA) breaches at registration time and define "
        "monitoring KPIs.",
    ])

    # 3 ---------------------------------------------------------------------
    rp.h1("3. Literature Review")
    rp.para("Text categorisation. Representing documents with TF-IDF weights [3] and training a linear classifier "
            "remains a strong baseline for short texts. Joachims showed that support vector machines suit the "
            "high-dimensional, sparse feature spaces of text [4], and Rennie et al. proposed Complement Naive "
            "Bayes to correct the bias of standard Naive Bayes on imbalanced classes [5]. Ensemble methods such as "
            "random forests [6] and gradient boosting [7] are widely used for tabular prediction. Character "
            "n-grams make models robust to spelling variation, which is important because Indian users frequently "
            "mix Hindi and English in Roman script [8].")
    rp.para("Calibration and selective prediction. SVM scores are not probabilities; Platt scaling maps them to "
            "calibrated probabilities [9]. Selective classification lets a model abstain on low-confidence inputs, "
            "trading coverage for accuracy [10], which matches a human-in-the-loop routing desk.")
    rp.para("Cost-sensitive decisions. When misclassification costs are unequal, the optimal decision is the "
            "class with minimum expected cost under the predicted class probabilities rather than the most "
            "probable class [11]. Ordinal agreement can be measured with weighted kappa [12]. Priority queues "
            "from queueing theory [13] explain why ordering by urgency shortens waits for critical jobs.")
    rp.para("Multimodal extraction and analytics. OCR engines such as Tesseract [14], image classifiers such as "
            "ResNet [15] and speech recognisers such as Whisper [16] turn images and audio into text or tags "
            "that a text pipeline can use. Count data are modelled with Poisson generalised linear models [17]. "
            "Clustering quality is measured with the adjusted Rand index [18], feature relevance with permutation "
            "importance [19], and opinion mining techniques are surveyed in [20]. scikit-learn [21] implements "
            "most of these methods.")

    # 4 ---------------------------------------------------------------------
    rp.h1("4. Existing System")
    rp.para("In the present workflow described in SIH1711 [2], a complaint is registered, acknowledged with an ID "
            "and categorised by staff or by the passenger choosing from drop-down menus, then forwarded to a "
            "department. Its main limitations are:")
    rp.bullets([
        ("Manual categorisation: ", "slow, and wrong categories cause re-routing between departments."),
        ("First-come-first-served handling: ", "a medical emergency can wait behind dozens of minor complaints."),
        ("Unused media: ", "photos, videos and audio are viewed manually; their text and metadata are not extracted."),
        ("Duplicate tickets: ", "many passengers in one coach report the same AC or water failure separately."),
        ("Reactive maintenance: ", "recurring faults on particular trains are not predicted from complaint trends."),
        ("Limited feedback analytics: ", "free-text feedback is not systematically analysed."),
    ])

    # 5 ---------------------------------------------------------------------
    rp.h1("5. Proposed Solution")
    rp.para("The proposed system, named RailSaathi AI (Fig. 1), is an analytical pipeline that sits behind the existing Rail Madad "
            "front end. Every complaint, whatever its medium, is converted into a single fused document: typed "
            "text or an ASR transcript, OCR text from attached images, tags from an image classifier, and "
            "structured metadata (channel, media type, train type, time). Three models act on this document: a "
            "calibrated categoriser that decides the department, a cost-sensitive urgency model that decides the "
            "priority, and a grouping step that merges duplicate reports. An analytics layer forecasts "
            "maintenance hotspots, predicts SLA breaches and tracks sentiment.")
    rp.figure("fig01_architecture", "Architecture of RailSaathi AI, the proposed AI-assisted Rail Madad complaint pipeline.", width=5.8)
    rp.para("The novel elements, compared with a plain text classifier, are: (i) multimodal late-fusion of text, "
            "OCR and vision tags in one sparse feature space; (ii) confidence-aware routing with a tunable accuracy "
            "target; (iii) a minimum-expected-cost severity rule evaluated with a discrete-event queue simulation; "
            "(iv) graph-based incident grouping; and (v) Poisson forecasting that turns complaints into predictive "
            "maintenance signals.")

    # 6 ---------------------------------------------------------------------
    rp.h1("6. Methodology")
    rp.h2("6.1 Dataset")
    rp.para("Real Rail Madad records are not publicly released, so a synthetic dataset was generated with a "
            "purpose-built simulator (src/generate_dataset.py). The 13 categories and their departments follow the "
            "Rail Madad portal [1]. The simulator is process-driven rather than random: every train carries a "
            "latent condition score per category that evolves as an AR(1) process, eight trains deteriorate "
            "steadily, older rakes produce more mechanical and electrical faults, and seasonal multipliers raise "
            "AC and water complaints in summer, leakage in the monsoon, fog-related delays in winter and crowding "
            "in festival months. Weekly complaint counts are drawn from a Poisson distribution and each complaint "
            "is then written in natural language.")
    rp.para(f"To keep the task realistic, {pct(ov['share_hinglish'])} of complaints are in Hinglish, about 30% "
            "contain typing errors or SMS spellings, 15% mention a second minor issue, 4% are vague, and 3% carry "
            "a deliberately wrong label (annotation noise). Photos and videos come with OCR text, imperfect vision "
            "tags (correct about 70% of the time, mimicking a CNN), EXIF capture time and GPS, while audio "
            "complaints arrive as noisy ASR transcripts. Some faults are reported by 2 to 5 passengers, giving "
            "ground-truth incident IDs. Resolution time follows a log-normal model driven by department, severity, "
            "zone efficiency, night-time and evidence; feedback ratings fall as resolution time exceeds the SLA.")
    schema = pd.DataFrame([
        ["Complaint", "complaint_text, channel, media_type, timestamp, is_hinglish"],
        ["Media-derived", "ocr_text, vision_tags, exif_capture_lag_min, gps_lat, gps_lon"],
        ["Journey", "train_no, train_type, zone, train_age_years, coach, pnr, nearest_station"],
        ["Labels", "category (13), department (7), severity (4), incident_id"],
        ["Process", "resolution_hours, sla_hours, sla_breached, feedback_rating, feedback_text"],
    ], columns=["Group", "Attributes"])
    rp.table(schema, "Schema of the synthetic Rail Madad dataset.", col_widths=[1.3, 4.9], index=False)
    rp.para(f"The final dataset holds {int(float(ov['complaints'])):,} complaints ({int(float(ov['incidents'])):,} "
            f"distinct incidents) on {int(float(ov['trains']))} trains from {ov['period'].replace(' to ', ' to ')}. "
            f"A time-based split is used: complaints before 1 April 2026 ({int(float(ov['train_rows'])):,}) for "
            f"training and the last six months ({int(float(ov['test_rows'])):,}) for testing, so every model is "
            "evaluated on future data exactly as it would be deployed.")
    rp.figure("fig02_category_severity", "Complaint volume by category, split by severity level.", width=4.5)

    rp.h2("6.2 Pre-processing and Feature Engineering")
    rp.para("Text is lower-cased; handles, PNRs and train numbers are replaced by placeholder tokens so that the "
            "classifier cannot memorise them; a compact (about 50-word) Hinglish lexicon (for example pani to water, ganda to "
            "dirty) and an SMS dictionary normalise code-mixed and abbreviated words. OCR tokens are prefixed "
            "with ocr_ and vision tags with img_, and all are concatenated into one document. Features are word "
            "1-2-gram and character 2-5-gram TF-IDF vectors [3]; the urgency model also receives one-hot metadata.")

    rp.h2("6.3 Models and Evaluation")
    rp.bullets([
        ("Categorisation: ", "Complement Naive Bayes, Logistic Regression, a Platt-calibrated linear SVM and a "
         "Random Forest, compared on accuracy, macro-F1 and department-routing accuracy."),
        ("Routing: ", "a complaint is auto-assigned if the top calibrated probability exceeds a threshold chosen so "
         "that auto-routed complaints reach 97% accuracy; the rest go to a human desk."),
        ("Urgency: ", "a cost matrix C penalises under-triage quadratically (Critical predicted as Low costs 16, "
         "over-triage costs 1). The predicted class is argmin_k sum_j p(j|x) C[j,k] [11]."),
        ("Queue simulation: ", "test complaints arrive at their real timestamps at a single triage desk "
         "(mean handling 30 min, 92% utilisation), handled by FIFO, by AI priority or by true priority."),
        ("Extraction and grouping: ", "regular expressions for train number, coach and PNR; haversine nearest "
         "station from GPS; complaints on the same train within 6 h with the same predicted category are linked "
         "if they share the booked coach or exceed a cosine-similarity threshold, and connected components "
         "become incidents, scored with pairwise F1 and ARI [18]."),
        ("Forecasting: ", "weekly counts per train of five maintenance categories, predicted with lag, rolling, "
         "trend and seasonal features using a Poisson GLM [17] and Poisson-loss gradient boosting [7]."),
        ("Sentiment and SLA: ", "a TF-IDF logistic-regression sentiment model trained on feedback text with labels "
         "from star ratings, and a gradient-boosting SLA-breach classifier explained by permutation importance [19]."),
    ])

    # 7 ---------------------------------------------------------------------
    rp.h1("7. Implementation")
    rp.para("The project is written in Python 3.11 using pandas, NumPy, SciPy, scikit-learn [21] and Matplotlib. "
            "Each stage is a separate module in src/, and run_all.py executes the full pipeline (data generation, "
            "EDA, modelling, figures and result tables) in under two minutes on a laptop. All random processes are "
            "seeded, so every number in this report can be reproduced. Results are written as CSV/JSON files to "
            "results/ and figures as 300-dpi PNG files to figures/. The source code is available at " + GITHUB + ".")
    impl = pd.DataFrame([
        ["generate_dataset.py", "Process-driven multimodal complaint simulator"],
        ["common.py", "Paths, plotting style, Hinglish/SMS normalisation, multimodal fusion"],
        ["eda.py", "Descriptive statistics, seasonality and operational analysis"],
        ["categorize.py", "Model comparison, ablation, confusion matrix, confidence routing"],
        ["urgency.py", "Cost-sensitive severity prediction and queue simulation"],
        ["extraction_dedup.py", "Entity extraction, GPS location, duplicate-incident grouping"],
        ["predictive_maintenance.py", "Weekly Poisson forecasting and hotspot detection"],
        ["sentiment_sla.py", "Sentiment model, SLA-breach risk model, KPI monitoring"],
    ], columns=["Module", "Responsibility"])
    rp.table(impl, "Implementation modules.", col_widths=[2.1, 4.1], index=False)
    rp.para("Exploratory analysis confirmed that the simulated operational patterns are visible in the data. "
            "Fig. 3 shows the relative monthly share of each category: electrical and water complaints peak in "
            "April-June, maintenance complaints in the monsoon, punctuality complaints in the foggy months of "
            "December-January and security complaints around festivals. The overall SLA breach rate is "
            f"{pct(ov['sla_breach_rate'])}, ranging from 19.0% in the best zone to 43.6% in the worst.")
    rp.figure("fig03_seasonality", "Seasonality of complaint categories (share relative to the category average).", width=5.3)

    # 8 ---------------------------------------------------------------------
    rp.h1("8. Results")
    rp.h2("8.1 Complaint Categorisation and Routing")
    t = comp[["Accuracy", "Macro-F1", "Dept. routing acc.", "Train time (s)"]].copy()
    for c in ["Accuracy", "Macro-F1", "Dept. routing acc."]:
        t[c] = (100 * t[c]).map("{:.2f}".format)
    t["Train time (s)"] = t["Train time (s)"].map("{:.1f}".format)
    t.columns = ["Accuracy (%)", "Macro-F1 (%)", "Dept. routing (%)", "Train time (s)"]
    t.index = [i + "*" if i == "Linear SVM (calibrated)" else i for i in t.index]
    t.index.name = "Model"
    rp.table(t, "Categorisation performance on the six-month test set (* best model).",
             col_widths=[2.2, 1.0, 1.0, 1.2, 1.0])
    rp.para(f"The calibrated linear SVM is the best model with {pct(best['Accuracy'], 2)} accuracy and "
            f"{pct(best['Macro-F1'], 2)} macro-F1, and it sends {pct(best['Dept. routing acc.'], 2)} of complaints "
            "to the correct department. It needs about 6 s to train and well under 0.1 ms per complaint to "
            "predict, so it can classify a complaint the moment it is lodged. Fig. 4 shows that most remaining "
            "errors involve the two smallest classes, Divyangjan Facilities and Women Safety, whose complaints "
            "share vocabulary with other classes.")
    rp.figure("fig06_confusion_matrix", "Normalised confusion matrix of the calibrated linear SVM.", width=3.9)
    a = abl.copy()
    a = (100 * a).round(2)
    a.index.name = "Input streams (linear SVM)"
    rp.table(a, "Ablation study: contribution of each input stream (accuracy and macro-F1 in %).",
             col_widths=[2.4, 0.95, 0.95, 1.0, 0.95], font=9)
    rp.para("The ablation (Table 4) shows that the Hinglish lexicon and character n-grams add only small gains, "
            "because character n-grams already match spelling variants. OCR text and especially vision tags "
            "help most: with full fusion, accuracy on photo/video complaints rises from "
            f"{abl.iloc[0]['Accuracy (photo/video)'] * 100:.2f}% to {abl.iloc[-1]['Accuracy (photo/video)'] * 100:.2f}%, "
            "showing that the media passengers attach carry useful category information. Across sub-groups the "
            f"model stays robust: {rob.loc['Audio (ASR transcript)', 'Accuracy'] * 100:.1f}% on audio transcripts and "
            f"{rob.loc['Hinglish', 'Accuracy'] * 100:.1f}% on Hinglish complaints.")
    rp.para(f"For confidence-aware routing (Fig. 5), a threshold of {routing['threshold']:.2f} auto-routes "
            f"{pct(routing['auto_routed_share'])} of complaints with {pct(routing['auto_routed_accuracy'])} "
            f"accuracy. Only {pct(routing['human_review_share'])} go to the human desk, and on these the model "
            f"would have been right just {pct(routing['human_review_accuracy_of_model'])} of the time, which "
            "confirms that the abstention targets genuinely ambiguous complaints.")
    rp.figure("fig07_routing_tradeoff", "Coverage versus accuracy of confidence-aware routing.", width=3.7)

    rp.h2("8.2 Urgency Detection and Prioritisation")
    u = urg[["Accuracy", "Quadratic kappa", "Critical recall", "Critical+High recall", "Under-triage rate", "Mean cost"]].copy()
    for c in u.columns[:-1]:
        u[c] = (100 * u[c]).map("{:.1f}".format) if c != "Quadratic kappa" else u[c].map("{:.3f}".format)
    u["Mean cost"] = u["Mean cost"].map("{:.3f}".format)
    u.columns = ["Acc. (%)", "QW kappa", "Crit. recall (%)", "Crit.+High recall (%)", "Under-triage (%)", "Mean cost"]
    u.index = [i.replace(" (proposed)", "*") for i in u.index]
    u.index.name = "Method"
    rp.table(u, "Severity prediction on the test set (* proposed). Lower cost and under-triage are better.",
             col_widths=[2.0, 0.6, 0.65, 0.75, 0.85, 0.8, 0.65], font=9)
    rp.para(f"Plain arg-max logistic regression is the most accurate learner after the SVM, but it misses "
            f"{100 - lr_u['Critical recall'] * 100:.1f}% of critical cases. The minimum-expected-cost rule accepts "
            f"a lower overall accuracy ({pct(prop['Accuracy'])}) in exchange for higher critical recall "
            f"({pct(prop['Critical recall'])}), Critical+High recall of {pct(prop['Critical+High recall'])} and an "
            f"under-triage rate of only {pct(prop['Under-triage rate'])}, four times lower than arg-max. It has the "
            "lowest mean cost of all methods. A keyword-only rule system performs poorly, and adding a "
            "safety-keyword override on top of the model gave no further benefit, because the learned model "
            "already captures those cues.")
    rp.para(f"The queue simulation (Fig. 6b) shows the operational impact: with FIFO handling, critical "
            f"complaints wait {qp.loc['Critical', 'FIFO (current)']:.0f} min on average; with AI priority this falls "
            f"to {qp.loc['Critical', 'AI priority (proposed)']:.0f} min, within 2 min of the ideal "
            f"({qp.loc['Critical', 'Oracle priority (upper bound)']:.0f} min). High-severity waits fall from "
            f"{qp.loc['High', 'FIFO (current)']:.0f} to {qp.loc['High', 'AI priority (proposed)']:.0f} min. The "
            "price is a longer wait for low-severity complaints, which have the most relaxed (48 h) SLA.")
    rp.figure("fig09_urgency_queue", "(a) Effect of the cost-sensitive decision rule; (b) mean waiting time by "
              "severity under three queueing policies.", width=5.8)

    rp.h2("8.3 Data Extraction and Duplicate Grouping")
    e = ext[ext["Source"] == "Text + OCR"].set_index("Field")
    e0 = ext[ext["Source"] == "Complaint text only"].set_index("Field")
    rp.para(f"The extractors are precise (train number {e.loc['train_no', 'Precision (%)']:.1f}%, coach "
            f"{e.loc['coach', 'Precision (%)']:.1f}%, PNR {e.loc['pnr', 'Precision (%)']:.1f}% precision) and "
            "recover almost every value that is actually mentioned. Reading the attached images with OCR raises "
            f"the share of complaints with a recovered train number from {e0.loc['train_no', 'Recall (%)']:.1f}% "
            f"to {e.loc['train_no', 'Recall (%)']:.1f}% and coach from {e0.loc['coach', 'Recall (%)']:.1f}% to "
            f"{e.loc['coach', 'Recall (%)']:.1f}%. EXIF GPS was present in {pct(gps['share_with_gps'])} of "
            f"complaints and identified the nearest station correctly in {pct(gps['nearest_station_accuracy'])} of them.")
    rp.para(f"For duplicate grouping (Fig. 7), text similarity alone reaches at best a pairwise F1 of "
            f"{dup_txt['pair_F1']:.3f}, because different passengers describe the same fault in different words. "
            f"Adding the booked coach from the PNR record raises pairwise F1 to {dup_best['pair_F1']:.3f} "
            f"(ARI {dup_best['ARI']:.3f}) and reduces {int(dup_best['tickets_before']):,} complaints to "
            f"{int(dup_best['tickets_after_grouping']):,} tickets, close to the "
            f"{int(dup_best['true_incidents']):,} true incidents. This removes about "
            f"{100 * (1 - dup_best['tickets_after_grouping'] / dup_best['tickets_before']):.0f}% of redundant work.")
    rp.figure("fig10_duplicate_grouping", "Pairwise precision, recall and F1 of duplicate grouping versus the "
              "similarity threshold.", width=3.9)

    rp.h2("8.4 Predictive Maintenance")
    m = mf.copy().round(3)
    m.index = [i + "*" if i == "Gradient Boosting (Poisson)" else i for i in m.index]
    m.index.name = "Model"
    rp.table(m, "Weekly maintenance-complaint forecasting per train (test period, * proposed).",
             col_widths=[2.4, 1.0, 1.0, 1.4])
    rp.para("Poisson-loss gradient boosting gives the lowest Poisson deviance, the proper error measure for "
            "counts, and cuts the deviance of the naive last-week forecast by about 65%. The 4-week moving average "
            "is competitive on MAE, which is expected for noisy weekly counts, but it cannot use season or train "
            "age. More importantly, ranking trains by forecast complaints flags hotspots: of the eight trains "
            f"flagged at the start of the test period, {hot['precision_at_k_deteriorating'] * 100:.0f}% are "
            f"truly deteriorating trains, against a random rate of {hot['random_baseline'] * 100:.1f}%, and "
            f"{hot['overlap_with_actual_future_top_k'] * 100:.1f}% are among the eight trains that actually "
            "produced the most complaints later. These rakes can be sent for inspection before failures pile up.")
    rp.figure("fig11_maintenance_forecast", "(a) Network-level weekly forecast; (b) history and forecast of the "
              "highest-risk train.", width=5.8)

    rp.h2("8.5 Sentiment, SLA Risk and Performance Monitoring")
    rp.para(f"The sentiment model classifies feedback with {pct(sent['accuracy'])} accuracy (macro-F1 "
            f"{sent['macro_F1']:.3f}) and recalls {pct(sent['negative_recall'])} of negative feedback. Negative "
            "sentiment is strongly tied to service level: within SLA only 5-8% of feedback is negative in every "
            "department, while after an SLA breach it rises to 31-42%, highest for Medical "
            f"({neg.loc['Medical', 'SLA breached'] * 100:.1f}%). Monthly breach rate and negative-feedback share are "
            f"correlated (r = {mon['monthly_corr_breach_vs_negative']:.2f}), so both can serve as KPIs. The "
            f"SLA-breach model predicts breaches at registration with ROC-AUC "
            f"{sla.loc['Logistic Regression', 'ROC-AUC']:.3f} (logistic regression) and "
            f"{sla.loc['Gradient Boosting', 'ROC-AUC']:.3f} (gradient boosting); the top 20% riskiest "
            f"complaints contain {sla.loc['Gradient Boosting', 'Breaches caught in top-20% risk'] * 100:.1f}% of "
            "all breaches. Severity, category, zone and night-time registration are the main drivers (Fig. 9a).")
    rp.figure("fig12_sla_sentiment", "(a) Permutation importance of SLA-breach drivers; (b) monthly KPI trend.", width=5.8)

    # 9 ---------------------------------------------------------------------
    rp.h1("9. Discussion")
    rp.para("The results support each objective, and Table 7 links them. Three observations stand out. First, "
            "multimodal fusion matters most where the existing system struggles: photo and video complaints. "
            "Second, accuracy is the wrong target for urgency. Accepting a few points of accuracy loss to reduce "
            "under-triage four-fold is the right trade-off when the cost of a missed emergency is far greater than "
            "the cost of a false alarm. The queue simulation turns this into waiting minutes that managers can "
            "understand. Third, structured metadata (PNR coach, GPS, train ID) often beats sophisticated text "
            "similarity, as seen in duplicate grouping, so integrating the model with railway databases is as "
            "important as the model itself.")
    summary = pd.DataFrame([
        ["Categorisation and routing", f"{pct(best['Accuracy'])} accuracy; {pct(routing['auto_routed_share'])} auto-routed at {pct(routing['auto_routed_accuracy'])}"],
        ["Multimodal analysis", f"Photo/video accuracy {abl.iloc[0]['Accuracy (photo/video)'] * 100:.1f}% to {abl.iloc[-1]['Accuracy (photo/video)'] * 100:.1f}% with OCR + vision tags"],
        ["Urgency detection", f"Critical recall {pct(prop['Critical recall'])}; under-triage {pct(prop['Under-triage rate'])}"],
        ["Prioritisation", f"Critical wait {qs['critical_wait_fifo']:.0f} to {qs['critical_wait_ai']:.0f} min"],
        ["Data extraction", f"Train number recovered +{e.loc['train_no', 'Recall (%)'] - e0.loc['train_no', 'Recall (%)']:.1f} pts with OCR; GPS station {pct(gps['nearest_station_accuracy'])}"],
        ["Duplicate grouping", f"Pairwise F1 {dup_best['pair_F1']:.3f}; {100 * (1 - dup_best['tickets_after_grouping'] / dup_best['tickets_before']):.0f}% fewer tickets"],
        ["Predictive maintenance", f"Deviance -65% vs naive; {hot['precision_at_k_deteriorating'] * 100:.0f}% precision@8 vs {hot['random_baseline'] * 100:.0f}% random"],
        ["Feedback and monitoring", f"Sentiment accuracy {pct(sent['accuracy'])}; SLA-breach AUC {sla.loc['Logistic Regression', 'ROC-AUC']:.2f}"],
    ], columns=["Objective", "Key result"])
    rp.table(summary, "Summary of results against the objectives.", col_widths=[2.0, 4.2], index=False)
    rp.para("Limitations. The dataset is synthetic. Its patterns are realistic by design, but the absolute "
            "scores, especially for duplicate grouping, where reports of one incident always share a coach, "
            "would be lower on real data. The vision tags and ASR transcripts are simulated outputs rather than "
            "the outputs of trained image and speech models. The SLA model reaches only moderate AUC because "
            "much of the resolution-time variation is random by construction, as it is in real operations.")

    # 10 --------------------------------------------------------------------
    rp.h1("10. Conclusion")
    rp.para("This project designed and evaluated RailSaathi AI, a complete data-science pipeline for SIH1711 that turns raw, "
            "multimodal passenger complaints into routed, prioritised and de-duplicated tickets, and turns the "
            "complaint stream into maintenance and service-quality intelligence. A calibrated linear SVM on fused "
            "text, OCR and image tags categorised 95% of complaints correctly. Confidence-aware routing automated "
            "almost 98% of assignments at 97% accuracy. A cost-sensitive urgency rule cut the average wait for "
            "critical complaints in a simulated triage desk from about nine hours to about half an hour. Graph-based "
            "grouping removed 18% of redundant tickets, and Poisson forecasting identified deteriorating trains "
            "nearly four times better than chance. All components are lightweight, fast and reproducible, so they "
            "could be added to the existing Rail Madad back end without changing the passenger interface.")

    # 11 --------------------------------------------------------------------
    rp.h1("11. Future Scope")
    rp.bullets([
        "Validate on real, anonymised Rail Madad records obtained from the Ministry of Railways.",
        "Replace simulated vision tags with a fine-tuned CNN or vision transformer on coach images, and use a "
        "speech model such as Whisper [16] on real voice complaints.",
        "Use multilingual transformer embeddings for regional languages beyond Hindi and English.",
        "Add a chatbot that asks for missing details (PNR, coach) when the extractors find none.",
        "Link predictive-maintenance alerts with depot maintenance schedules and optimise staff allocation "
        "across zones from the forecasts.",
    ])

    # references ------------------------------------------------------------
    rp.h1("References")
    refs = [
        'Ministry of Railways, Government of India, "Rail Madad - For inquiry, assistance and grievance redressal." '
        "[Online]. Available: https://railmadad.indianrailways.gov.in",
        'Ministry of Education\'s Innovation Cell, "Smart India Hackathon 2024 - Problem statement SIH1711: Enhancing '
        'Rail Madad with AI-powered complaint management," 2024.',
        'G. Salton and C. Buckley, "Term-weighting approaches in automatic text retrieval," Information Processing & '
        "Management, vol. 24, no. 5, pp. 513-523, 1988.",
        'T. Joachims, "Text categorization with support vector machines: Learning with many relevant features," in '
        "Proc. 10th European Conf. Machine Learning (ECML), 1998, pp. 137-142.",
        'J. D. M. Rennie, L. Shih, J. Teevan, and D. R. Karger, "Tackling the poor assumptions of naive Bayes text '
        'classifiers," in Proc. 20th Int. Conf. Machine Learning (ICML), 2003, pp. 616-623.',
        'L. Breiman, "Random forests," Machine Learning, vol. 45, no. 1, pp. 5-32, 2001.',
        'J. H. Friedman, "Greedy function approximation: A gradient boosting machine," The Annals of Statistics, '
        "vol. 29, no. 5, pp. 1189-1232, 2001.",
        'K. Bali, J. Sharma, M. Choudhury, and Y. Vyas, "\'I am borrowing ya mixing?\' An analysis of English-Hindi '
        'code mixing in Facebook," in Proc. First Workshop on Computational Approaches to Code Switching, 2014, '
        "pp. 116-126.",
        'J. Platt, "Probabilistic outputs for support vector machines and comparisons to regularized likelihood '
        'methods," in Advances in Large Margin Classifiers, MIT Press, 1999, pp. 61-74.',
        'Y. Geifman and R. El-Yaniv, "Selective classification for deep neural networks," in Advances in Neural '
        "Information Processing Systems 30 (NeurIPS), 2017.",
        'C. Elkan, "The foundations of cost-sensitive learning," in Proc. 17th Int. Joint Conf. Artificial '
        "Intelligence (IJCAI), 2001, pp. 973-978.",
        'J. Cohen, "Weighted kappa: Nominal scale agreement with provision for scaled disagreement or partial '
        'credit," Psychological Bulletin, vol. 70, no. 4, pp. 213-220, 1968.',
        "L. Kleinrock, Queueing Systems, Volume 1: Theory. New York, NY, USA: Wiley, 1975.",
        'R. Smith, "An overview of the Tesseract OCR engine," in Proc. 9th Int. Conf. Document Analysis and '
        "Recognition (ICDAR), 2007, pp. 629-633.",
        'K. He, X. Zhang, S. Ren, and J. Sun, "Deep residual learning for image recognition," in Proc. IEEE Conf. '
        "Computer Vision and Pattern Recognition (CVPR), 2016, pp. 770-778.",
        'A. Radford, J. W. Kim, T. Xu, G. Brockman, C. McLeavey, and I. Sutskever, "Robust speech recognition via '
        'large-scale weak supervision," in Proc. 40th Int. Conf. Machine Learning (ICML), 2023, pp. 28492-28518.',
        "P. McCullagh and J. A. Nelder, Generalized Linear Models, 2nd ed. London, U.K.: Chapman & Hall, 1989.",
        'L. Hubert and P. Arabie, "Comparing partitions," Journal of Classification, vol. 2, no. 1, pp. 193-218, 1985.',
        'A. Altmann, L. Tolosi, O. Sander, and T. Lengauer, "Permutation importance: A corrected feature importance '
        'measure," Bioinformatics, vol. 26, no. 10, pp. 1340-1347, 2010.',
        'B. Pang and L. Lee, "Opinion mining and sentiment analysis," Foundations and Trends in Information '
        "Retrieval, vol. 2, no. 1-2, pp. 1-135, 2008.",
        'F. Pedregosa et al., "Scikit-learn: Machine learning in Python," Journal of Machine Learning Research, '
        "vol. 12, pp. 2825-2830, 2011.",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.45)
        p.paragraph_format.first_line_indent = Inches(-0.45)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(f"[{i}]\t{ref}"), size=11)

    doc.save(OUT)
    print(f"Saved {OUT}  ({rp.fig_no} figures, {rp.tab_no} tables, {len(refs)} references)")


if __name__ == "__main__":
    build()
